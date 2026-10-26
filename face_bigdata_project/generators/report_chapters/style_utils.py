"""
Các tiện ích định dạng văn bản chuẩn học thuật cho Báo cáo Đồ án
"""

import os
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import pandas as pd

def setup_document_styles(doc):
    """Cấu hình kích thước A4 và lề chuẩn quy định."""
    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(3.0)
        section.right_margin = Cm(2.0)

    # Cấu hình Style Normal
    style_normal = doc.styles['Normal']
    font = style_normal.font
    font.name = 'Times New Roman'
    font.size = Pt(13)
    font.color.rgb = RGBColor(30, 41, 59)

def add_p(doc, text="", space_before=0, space_after=6, line_spacing=1.35, align=WD_ALIGN_PARAGRAPH.JUSTIFY, bold=False, italic=False, font_size=13, color=None):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing

    if text:
        run = p.add_run(text)
        run.bold = bold
        run.italic = italic
        run.font.name = 'Times New Roman'
        run.font.size = Pt(font_size)
        if color:
            run.font.color.rgb = color
    return p

def add_h1(doc, text):
    """Heading 1: Ngắt trang, 16pt, Đậm, In hoa."""
    doc.add_page_break()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text.upper())
    run.bold = True
    run.font.name = 'Times New Roman'
    run.font.size = Pt(16)
    run.font.color.rgb = RGBColor(14, 116, 144)
    return p

def add_h2(doc, text):
    """Heading 2: 14pt, Đậm."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.bold = True
    run.font.name = 'Times New Roman'
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(15, 23, 42)
    return p

def add_h3(doc, text):
    """Heading 3: 13pt, Đậm nghiêng."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.bold = True
    run.italic = True
    run.font.name = 'Times New Roman'
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor(30, 41, 59)
    return p

def add_formula(doc, formula_text, formula_number=""):
    """Định dạng khối công thức toán học chuyên nghiệp."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.left_indent = Inches(0.5)

    run = p.add_run(formula_text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    run.italic = True
    run.bold = True
    run.font.color.rgb = RGBColor(15, 23, 42)

    if formula_number:
        p_num = doc.add_paragraph()
        p_num.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p_num.paragraph_format.space_before = Pt(0)
        p_num.paragraph_format.space_after = Pt(6)
        r_num = p_num.add_run(f"({formula_number})")
        r_num.font.name = 'Times New Roman'
        r_num.font.size = Pt(11)
        r_num.italic = True

def add_callout_box(doc, title, text):
    """Tạo hộp ghi chú / định lý nổi bật."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F1F5F9"/>')
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="24" w:space="0" w:color="0E7490"/><w:top w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/></w:tcBorders>')
    tcPr.append(shd)
    tcPr.append(borders)

    p0 = cell.paragraphs[0]
    p0.paragraph_format.space_before = Pt(4)
    p0.paragraph_format.space_after = Pt(2)
    r0 = p0.add_run(title)
    r0.bold = True
    r0.font.name = 'Times New Roman'
    r0.font.size = Pt(11.5)
    r0.font.color.rgb = RGBColor(14, 116, 144)

    p1 = cell.add_paragraph()
    p1.paragraph_format.space_before = Pt(2)
    p1.paragraph_format.space_after = Pt(4)
    p1.paragraph_format.line_spacing = 1.25
    r1 = p1.add_run(text)
    r1.font.name = 'Times New Roman'
    r1.font.size = Pt(11)
    r1.italic = True
    r1.font.color.rgb = RGBColor(51, 65, 85)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_figure(doc, figure_filename, caption_text, figures_dir, width_inches=5.8):
    image_path = os.path.join(figures_dir, figure_filename)
    if os.path.exists(image_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(4)
        run_img = p_img.add_run()
        run_img.add_picture(image_path, width=Inches(width_inches))

        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(2)
        p_cap.paragraph_format.space_after = Pt(10)
        run_cap = p_cap.add_run(caption_text)
        run_cap.italic = True
        run_cap.font.name = 'Times New Roman'
        run_cap.font.size = Pt(11)
        run_cap.font.color.rgb = RGBColor(71, 85, 105)

def add_table_from_csv(doc, csv_filename, caption_text, tables_dir):
    csv_path = os.path.join(tables_dir, csv_filename)
    if not os.path.exists(csv_path):
        return

    df = pd.read_csv(csv_path)

    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_cap.paragraph_format.space_before = Pt(10)
    p_cap.paragraph_format.space_after = Pt(4)
    run_cap = p_cap.add_run(caption_text)
    run_cap.bold = True
    run_cap.italic = True
    run_cap.font.name = 'Times New Roman'
    run_cap.font.size = Pt(11.5)
    run_cap.font.color.rgb = RGBColor(14, 116, 144)

    n_rows = len(df) + 1
    n_cols = len(df.columns)
    table = doc.add_table(rows=n_rows, cols=n_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    hdr_row = table.rows[0]
    for idx, col_name in enumerate(df.columns):
        cell = hdr_row.cells[idx]
        cell.text = str(col_name).replace("_", " ")
        set_cell_background(cell, "0E7490")
        set_cell_margins(cell, top=140, bottom=140, left=140, right=140)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.name = 'Times New Roman'
            r.font.size = Pt(10.5)
            r.font.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)

    for r_idx, row in df.iterrows():
        row_cells = table.rows[r_idx + 1].cells
        bg_color = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row):
            cell = row_cells[c_idx]
            cell.text = str(val)
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx > 0 else WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.font.name = 'Times New Roman'
                r.font.size = Pt(10.5)
                r.font.color.rgb = RGBColor(30, 41, 59)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)
